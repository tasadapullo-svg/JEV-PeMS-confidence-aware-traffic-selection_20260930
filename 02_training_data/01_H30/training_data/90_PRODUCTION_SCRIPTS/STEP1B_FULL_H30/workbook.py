"""Create the 25-sheet Step 1B results workbook with review-oriented formatting."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SHEETS = [
    '00_Summary','01_DataScope','02_FrozenDesign','03_ModelConfig','04_Accuracy_AllScopes',
    '05_PrimaryTest','06_QualityBins','07_Q0_Degraded','08_SeenUnseen','09_PeakPeriod',
    '10_WeekdayWeekend','11_FreewayResults','12_XGB_GRU_Paired','13_ErrorThresholds',
    '14_Latency','15_TrainingRuntime','16_GRU_EpochHistory','17_HA_Fallback',
    '18_ResourceUsage','19_QC','20_FileManifest','21_PaperReady_Table1',
    '22_PaperReady_Table2','23_PaperReady_Table3','24_Notes_Limitations',
]

def read_csv(path):
    if not path.exists(): return []
    with path.open(newline='',encoding='utf-8-sig') as f: return list(csv.DictReader(f))

def rows_from_dict(obj):
    return [{'Field':k,'Value':json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v} for k,v in obj.items()]

def table(ws, rows, title=None):
    if title: ws.sheet_properties.pageSetUpPr.fitToPage=True
    if not rows:
        ws.append(['No rows']);return
    fields=list(rows[0])
    ws.append(fields)
    for r in rows:
        values=[]
        for f in fields:
            v=r.get(f)
            if isinstance(v,str):
                try:
                    if v.strip() and v.strip() not in ('NaN','nan'):
                        v=float(v) if any(c in v for c in '.eE') else int(v)
                except (ValueError,OverflowError): pass
            values.append(v)
        ws.append(values)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        cell.border=Border(bottom=Side(style='thin',color='808080'))
    ws.row_dimensions[1].height=32
    for col in ws.columns:
        letter=get_column_letter(col[0].column)
        length=max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:250])
        ws.column_dimensions[letter].width=max(13,min(52,length+2))
        header=str(col[0].value).lower()
        for cell in list(col)[1:]:
            if isinstance(cell.value,float):
                cell.number_format='0.00' if any(x in header for x in ('percent','_pct','seconds','time')) else '0.0000'
            cell.alignment=Alignment(vertical='top')

def build(run_dir:Path, output:Path, run_id:str, status:str):
    data_scope=json.loads((run_dir/'01_EVALUATION_SCOPE/evaluation_scope.json').read_text(encoding='utf-8'))
    config=json.loads((run_dir/'00_CONFIG/Step1B_Full_H30_Config.json').read_text(encoding='utf-8'))
    lock=json.loads((run_dir/'00_CONFIG/model_lock.json').read_text(encoding='utf-8'))
    qc=json.loads((run_dir/'12_QC/full_h30_qc.json').read_text(encoding='utf-8'))
    resource=json.loads((run_dir/'11_RESOURCE_PROFILE/compute_profile.json').read_text(encoding='utf-8'))
    accuracy=read_csv(run_dir/'07_METRICS/full_h30_forecasting_results.csv')
    quality=read_csv(run_dir/'08_SUBGROUP_ANALYSIS/full_h30_quality_results.csv')
    q0=read_csv(run_dir/'08_SUBGROUP_ANALYSIS/q0_degraded_results.csv')
    subgroups=read_csv(run_dir/'08_SUBGROUP_ANALYSIS/subgroup_results.csv')
    paired=read_csv(run_dir/'09_PAIRED_COMPARISON/xgb_vs_gru_paired_error.csv')
    thresholds=read_csv(run_dir/'09_PAIRED_COMPARISON/error_thresholds.csv')
    latency=read_csv(run_dir/'10_LATENCY/latency_benchmark.csv')
    epochs=read_csv(run_dir/'05_GRU/gru_epoch_history.csv')
    fallback=read_csv(run_dir/'03_HISTORICAL_AVERAGE/historical_average_fallback_summary.csv')
    manifest=read_csv(run_dir/'15_MANIFEST/Step1B_Full_H30_Manifest.csv')
    primary=[r for r in accuracy if r.get('scope')=='PRIMARY_TEST']
    paired_primary=next((r for r in paired if r.get('scope')=='PRIMARY_TEST'),{})
    summary=dict(STEP='STEP 1B FULL H30',RUN_ID=run_id,RUN_DATE=config.get('run_date'),STATUS=status,
        HORIZON='H30',TARGET_POLICY='%Observed = 100',TRAIN_RANGE=config['splits']['train'],
        VALIDATION_RANGE=config['splits']['validation'],TEST_RANGE=config['splits']['test'],
        TRAIN_ROWS=data_scope['train_rows'],PRIMARY_VAL_ROWS=data_scope['primary_validation_rows'],
        PRIMARY_TEST_ROWS=data_scope['primary_test_rows'],FULL_TEST_ROWS=data_scope['full_test_rows'],
        UNSEEN_TEST_ROWS=data_scope['unseen_test_rows'],FAST_MODEL_CANDIDATE='XGBoost',
        STRONG_MODEL_CANDIDATE='GRU',GRU_better_pct=paired_primary.get('GRU_better_pct'),
        XGB_better_pct=paired_primary.get('XGB_better_pct'),mean_delta_AE=paired_primary.get('mean_delta_AE'),
        Leakage=qc.get('leakage_violations'),Duplicates=qc.get('duplicate_sample_id'),
        Missing=qc.get('missing_target'),Target_violations=qc.get('target100_violations'),
        Important_warnings='Smoke results are not paper results' if config.get('smoke') else 'No Jev or routing result in Step 1B')
    for r in primary:
        m=r['model'];summary[f'{m}_MAE']=r.get('MAE');summary[f'{m}_RMSE']=r.get('RMSE');summary[f'{m}_P95']=r.get('P95_AE')
    for m in ('XGBoost','GRU'):
        summary[f'{m}_E2E_latency_seconds']=next((r.get('median_seconds') for r in latency if r.get('model')==m and r.get('mode')=='END_TO_END'),None)
        summary[f'{m}_training_seconds']=resource.get(m,{}).get('training_seconds')
    notes=[{'Fact':x} for x in ['CPU-only benchmark','Station ID is not a predictive feature',
        'Primary evaluation uses Train-seen stations','Unseen stations retained as robustness cohort',
        'Target requires %Observed = 100','Q0 is degraded-input condition, not failure label',
        'No Jev in Step 1B','No routing label generated','H60 not included','Test not used for tuning',
        'Smoke metrics are only pipeline checks' if config.get('smoke') else 'Formal metrics are chronological H30 results']]
    paper1=[{k:r.get(k) for k in ('model','scope','N','stations','MAE','RMSE','Bias','P90_AE','P95_AE','P99_AE')} for r in accuracy]
    paper2=[{k:r.get(k) for k in ('model','quality_bin','N','MAE','RMSE','P95_AE')} for r in quality]
    latency_map={r['model']:r for r in latency if r.get('mode')=='END_TO_END'}
    xgb_cost=float(latency_map['XGBoost']['median_seconds']) if 'XGBoost' in latency_map else 0
    paper3=[]
    for r in primary:
        m=r['model'];l=latency_map.get(m,{})
        paper3.append(dict(Model=m,MAE=r.get('MAE'),RMSE=r.get('RMSE'),P95=r.get('P95_AE'),
            Training_Time=resource.get(m,{}).get('training_seconds'),End_to_End_Inference=l.get('median_seconds'),
            Samples_per_second=l.get('samples_per_second'),
            Relative_Compute_Cost=(float(l['median_seconds'])/xgb_cost if l and xgb_cost else None)))
    source={
        '00_Summary':rows_from_dict(summary),'01_DataScope':rows_from_dict(data_scope),
        '02_FrozenDesign':rows_from_dict(config),'03_ModelConfig':rows_from_dict(config['models'])+rows_from_dict(lock),
        '04_Accuracy_AllScopes':accuracy,'05_PrimaryTest':primary,'06_QualityBins':quality,
        '07_Q0_Degraded':q0,'08_SeenUnseen':[r for r in subgroups if r.get('dimension')=='seen_unseen'],
        '09_PeakPeriod':[r for r in subgroups if r.get('dimension')=='peak_period'],
        '10_WeekdayWeekend':[r for r in subgroups if r.get('dimension')=='weekday_weekend'],
        '11_FreewayResults':[r for r in subgroups if r.get('dimension')=='freeway'],
        '12_XGB_GRU_Paired':paired,'13_ErrorThresholds':thresholds,'14_Latency':latency,
        '15_TrainingRuntime':rows_from_dict(resource),'16_GRU_EpochHistory':epochs,
        '17_HA_Fallback':fallback,'18_ResourceUsage':rows_from_dict(resource),
        '19_QC':rows_from_dict(qc),'20_FileManifest':manifest,
        '21_PaperReady_Table1':paper1,'22_PaperReady_Table2':paper2,
        '23_PaperReady_Table3':paper3,'24_Notes_Limitations':notes,
    }
    wb=Workbook();wb.remove(wb.active)
    for name in SHEETS:table(wb.create_sheet(name),source[name],name)
    output.parent.mkdir(parents=True,exist_ok=True)
    wb.save(output)
    check=load_workbook(output,read_only=True,data_only=True)
    names=check.sheetnames;check.close()
    if names!=SHEETS:raise RuntimeError('XLSX sheet readback mismatch')
    return len(names)
